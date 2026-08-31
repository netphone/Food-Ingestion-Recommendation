"""Compare slide content of both files."""
from pptx import Presentation

def get_all_text(prs):
    result = []
    for i, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    slide_texts.append(t[:200])
        result.append((i+1, slide_texts))
    return result

f1 = '/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised (1).pptx'
f2 = '/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised_v5.pptx'

prs1 = Presentation(f1)
prs2 = Presentation(f2)

print(f'File 1: {len(prs1.slides)} slides')
print(f'File 2: {len(prs2.slides)} slides')
print()

t1 = get_all_text(prs1)
t2 = get_all_text(prs2)

max_slides = max(len(t1), len(t2))
identical = True

for i in range(max_slides):
    s1 = t1[i] if i < len(t1) else (i+1, [])
    s2 = t2[i] if i < len(t2) else (i+1, [])
    
    # Get first text as title
    title1 = s1[1][0].split('\n')[0][:60] if s1[1] else '[empty]'
    title2 = s2[1][0].split('\n')[0][:60] if s2[1] else '[empty]'
    
    # Compare all text
    texts1 = '\n'.join(s1[1])
    texts2 = '\n'.join(s2[1])
    
    match = texts1 == texts2
    if not match:
        identical = False
    
    status = 'OK' if match else 'DIFF'
    print(f'Slide {i+1:2d}: [{status}] {title1}')
    if not match:
        # Show differences
        for j in range(max(len(s1[1]), len(s2[1]))):
            t_a = s1[1][j][:100] if j < len(s1[1]) else '[missing]'
            t_b = s2[1][j][:100] if j < len(s2[1]) else '[missing]'
            if t_a != t_b:
                print(f'         F1: {t_a}')
                print(f'         F2: {t_b}')

print(f'\nIdentical: {identical}')
