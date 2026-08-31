from pptx import Presentation

prs = Presentation('/Users/linp24/Downloads/Quantative Stain Analytics_v2.pptx')
slide3 = prs.slides[2]

for shape in slide3.shapes:
    if ';193;' in shape.name and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if 'best linear' in run.text:
                    old = run.text
                    # Replace using substring that avoids quote issues
                    run.text = old.replace('identifies the mathematically proven ', 'yields the ')
                    # Remove smart quotes around BLUE phrase
                    run.text = run.text.replace('\u201c', '')
                    run.text = run.text.replace('\u201d', '')
                    print(f'Before: {old}')
                    print(f'After:  {run.text}')

prs.save('/Users/linp24/Downloads/Quantative Stain Analytics_v3.pptx')
print('\nSaved v3.')

# Verify
prs2 = Presentation('/Users/linp24/Downloads/Quantative Stain Analytics_v3.pptx')
for para in prs2.slides[2].shapes[2].text_frame.paragraphs:
    if 'quality metric' in para.text.lower():
        print(f'Verified: {para.text}')
