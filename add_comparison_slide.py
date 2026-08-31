"""
Insert a comparison slide after Slide 7 (second HPS slide) and before Slide 8 (HE_Quant).
Uses a table-based layout per Comment 3, with literature context.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INPUT = '/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised_v2.pptx'
OUTPUT = '/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised_v3.pptx'

prs = Presentation(INPUT)

# Find BLANK layout
blank_layout = None
for layout in prs.slide_layouts:
    if layout.name == 'BLANK':
        blank_layout = layout
        break

new_slide = prs.slides.add_slide(blank_layout)

# Move to index 7 (after Slide 7, before Slide 8)
slide_list = prs.slides._sldIdLst
ids = list(slide_list)
moved = ids[-1]
slide_list.remove(moved)
slide_list.insert(7, moved)

slide = prs.slides[7]

# --- Title ---
tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.0), Inches(0.45))
p = tx.text_frame.paragraphs[0]
r = p.add_run()
r.text = "Existing Methods: Strengths and Remaining Gap"
r.font.size = Pt(22)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

# --- Comparison Table ---
rows, cols = 6, 4
tbl_shape = slide.shapes.add_table(rows, cols, Inches(0.4), Inches(1.0), Inches(9.2), Inches(2.6))
table = tbl_shape.table

# Column widths
table.columns[0].width = Inches(2.0)
table.columns[1].width = Inches(2.2)
table.columns[2].width = Inches(2.2)
table.columns[3].width = Inches(2.8)

# Header style
header_bg = RGBColor(0x1A, 0x3C, 0x6E)
header_font = RGBColor(0xFF, 0xFF, 0xFF)

# Data
data = [
    ["Capability", "iQuant", "HemeProScore", "Remaining Need"],
    ["Calibration\nanchor",
     "Controlled HTX\ntiter series",
     "Reference cohort /\nimage",
     "Quantitative calibration\ntied to stain amount"],
    ["Stain\nseparation",
     "HTX-focused\n(no separation)",
     "H&E deconvolution\n(Ruifrok & Johnston)",
     "Calibrated, multi-stain\nseparation with quality metric"],
    ["Tissue\ncompartment",
     "Nuclear ROI only",
     "Whole-slide (global)",
     "Compartment-specific:\nnucleus + cytoplasm"],
    ["Scanner /\nprotocol",
     "Single protocol",
     "Relative normalization\nacross cohorts",
     "Absolute consistency\nacross scanners & protocols"],
    ["Quantification\ntype",
     "Absolute intensity\n(HTX only)",
     "Relative concentration\n(Vahadane / Macenko)",
     "Absolute OD-based\nquantification for both stains"],
]

for row_idx, row_data in enumerate(data):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = cell_text
        run.font.size = Pt(9)

        if row_idx == 0:
            # Header row
            from pptx.oxml.ns import qn
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            solidFill = tcPr.makeelement(qn('a:solidFill'), {})
            srgbClr = solidFill.makeelement(qn('a:srgbClr'), {'val': '1A3C6E'})
            solidFill.append(srgbClr)
            tcPr.append(solidFill)
            run.font.color.rgb = header_font
            run.font.bold = True
            run.font.size = Pt(10)
        elif col_idx == 0:
            # First column bold
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        elif col_idx == 3:
            # "Remaining Need" column in accent color
            run.font.color.rgb = RGBColor(0x1A, 0x73, 0xC8)
            run.font.bold = True
            run.font.size = Pt(9)
        else:
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# --- Bottom takeaway ---
tx_gap = slide.shapes.add_textbox(Inches(0.4), Inches(3.8), Inches(9.2), Inches(0.9))
tf = tx_gap.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
r1 = p1.add_run()
r1.text = "The Gap"
r1.font.size = Pt(13)
r1.font.bold = True
r1.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = (
    "Most stain-normalization approaches are fundamentally relative and reference-based "
    "(Macenko et al. 2009; Vahadane et al. 2016). They map stain appearance to a template "
    "but do not calibrate against an experimentally defined stain scale. "
    "A unified framework is needed for calibrated stain-vector estimation, compartment-specific "
    "quantification, and absolute consistency across imaging conditions."
)
r2.font.size = Pt(9)
r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
r2.font.italic = True

# --- Literature footnote ---
tx_ref = slide.shapes.add_textbox(Inches(0.4), Inches(4.8), Inches(9.2), Inches(0.4))
tf_ref = tx_ref.text_frame
tf_ref.word_wrap = True
p_ref = tf_ref.paragraphs[0]
r_ref = p_ref.add_run()
r_ref.text = (
    "References: Ruifrok & Johnston (2001) Anal Quant Cytol Histol; "
    "Macenko et al. (2009) IEEE ISBI; "
    "Vahadane et al. (2016) IEEE TMI"
)
r_ref.font.size = Pt(7)
r_ref.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# Save
prs.save(OUTPUT)
print(f"Saved: {OUTPUT}")

# Verify
prs2 = Presentation(OUTPUT)
print(f"Total slides: {len(prs2.slides)}")
for i in range(6, 10):
    s = prs2.slides[i]
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t and len(t) > 5:
                texts.append(t[:80])
    print(f"  Slide {i+1}: {texts[0] if texts else '[empty]'}")
