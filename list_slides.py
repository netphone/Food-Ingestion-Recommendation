from pptx import Presentation

prs = Presentation('/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised_v3.pptx')
print(f'Total: {len(prs.slides)} slides\n')

for i in range(len(prs.slides)):
    s = prs.slides[i]
    has_table = False
    has_image = False
    title = ''
    for shape in s.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t and len(t) > 3 and not title:
                title = t.split('\n')[0][:80]
        try:
            if shape.image:
                has_image = True
        except:
            pass
        if hasattr(shape, 'has_table') and shape.has_table:
            has_table = True
    extras = []
    if has_table:
        extras.append('TABLE')
    if has_image:
        extras.append('IMG')
    extra_str = f' [{", ".join(extras)}]' if extras else ''
    print(f'Slide {i+1:2d}: {title}{extra_str}')
