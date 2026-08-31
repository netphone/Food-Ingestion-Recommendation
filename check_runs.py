from pptx import Presentation

def find_in_groups(shape, kw):
    results = []
    if shape.has_text_frame and kw in shape.text_frame.text:
        results.append(shape)
    if shape.shape_type == 6:
        for c in shape.shapes:
            results.extend(find_in_groups(c, kw))
    return results

prs = Presentation('/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised.pptx')

print('=== Slide 2: runs with molecular ===')
for sh in prs.slides[1].shapes:
    for s in find_in_groups(sh, 'molecular'):
        for i, para in enumerate(s.text_frame.paragraphs):
            if para.text.strip():
                print(f'P{i}: ({len(para.runs)} runs)')
                for j, run in enumerate(para.runs):
                    print(f'  R{j}: {repr(run.text[:100])}')

print('\n=== Slide 3: runs with quality metric ===')
for sh in prs.slides[2].shapes:
    for s in find_in_groups(sh, 'quality metric'):
        for i, para in enumerate(s.text_frame.paragraphs):
            if para.text.strip():
                print(f'P{i}: ({len(para.runs)} runs)')
                for j, run in enumerate(para.runs):
                    print(f'  R{j}: {repr(run.text[:100])}')
