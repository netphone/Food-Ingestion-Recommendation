"""Apply comment fixes with correct run-level targeting."""
from pptx import Presentation

INPUT = '/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised.pptx'
OUTPUT = '/Users/linp24/Downloads/Quantative_Stain_Analytics_slide4_revised_v2.pptx'

def find_in_groups(shape, kw):
    results = []
    if shape.has_text_frame and kw in shape.text_frame.text:
        results.append(shape)
    if shape.shape_type == 6:
        for c in shape.shapes:
            results.extend(find_in_groups(c, kw))
    return results

prs = Presentation(INPUT)

# ============================================================
# Comment 1: Fix Slide 2
# ============================================================
for sh in prs.slides[1].shapes:
    for s in find_in_groups(sh, 'molecular'):
        for para in s.text_frame.paragraphs:
            for run in para.runs:
                # Fix 1a: P3 R2: "(molecular density)" -> stain-amount proxy language
                if '(molecular density)' in run.text:
                    run.text = run.text.replace(
                        '(molecular density)',
                        ', which reflects the local amount of absorbing dye,'
                    )
                    print('Fixed S2: molecular density')

                # Fix 1b: P5 R0: linearity claim
                if 'stain concentration is linearly proportional to light absorption' in run.text:
                    run.text = run.text.replace(
                        'ensuring that stain concentration is linearly proportional to light absorption.',
                        'so that optical density serves as an additive, stain-amount proxy. Exact molecular concentration requires additional calibration.'
                    )
                    print('Fixed S2: linearity claim')

# ============================================================
# Comment 2: Fix Slide 3
# ============================================================
for sh in prs.slides[2].shapes:
    for s in find_in_groups(sh, 'quality metric'):
        for para in s.text_frame.paragraphs:
            runs = list(para.runs)
            # Target P2: 5 runs
            # R0: 'A statistical quality metric '
            # R1: '(Qpar)'
            # R2: ' that yields the '
            # R3: 'best linear unbiased estimate (BLUE)'
            # R4: ' for stain vectors in Optical Density space.'
            if len(runs) >= 5:
                r0_text = runs[0].text
                r3_text = runs[3].text
                if 'quality metric' in r0_text and 'BLUE' in r3_text:
                    # Rewrite: "A statistical quality metric (Qpar) used to assess
                    # stain-separation quality and guide stain-vector calibration
                    # in Optical Density space."
                    runs[0].text = 'A statistical quality metric '
                    runs[1].text = '(Qpar)'
                    runs[2].text = ' used to assess '
                    runs[3].text = 'stain-separation quality'
                    runs[4].text = ' and guide stain-vector calibration in Optical Density space.'
                    print('Fixed S3: BLUE attribution removed, reworded')

# Save
prs.save(OUTPUT)
print(f'\nSaved: {OUTPUT}')

# Verify
prs2 = Presentation(OUTPUT)
print('\n=== Verify Slide 2 (relevant paragraphs) ===')
for sh in prs2.slides[1].shapes:
    for s in find_in_groups(sh, 'characterized'):
        for para in s.text_frame.paragraphs:
            t = para.text.strip()
            if 'characterized' in t or 'Beer-Lambert' in t or 'proxy' in t:
                print(f'  {t}')

print('\n=== Verify Slide 3 ===')
for sh in prs2.slides[2].shapes:
    for s in find_in_groups(sh, 'quality metric'):
        for para in s.text_frame.paragraphs:
            t = para.text.strip()
            if 'quality metric' in t.lower():
                print(f'  {t}')
