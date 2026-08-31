from pptx import Presentation

prs = Presentation('/Users/linp24/Downloads/Quantative Stain Analytics.pptx')

print(f'Slide width: {prs.slide_width}, height: {prs.slide_height}')
print(f'Total slides: {len(prs.slides)}\n')

print('=== Available Layouts ===')
for i, layout in enumerate(prs.slide_layouts):
    print(f'  [{i}] {layout.name}')
print()

for i, slide in enumerate(prs.slides):
    print(f'=== Slide {i+1} ===')
    print(f'Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        stype = str(shape.shape_type)
        print(f'  Shape: type={stype}, name="{shape.name}"')
        print(f'    pos=({shape.left},{shape.top}), size=({shape.width},{shape.height})')
        if shape.has_text_frame:
            for j, para in enumerate(shape.text_frame.paragraphs):
                text = para.text.strip()
                if text:
                    print(f'    Para[{j}]: "{text[:150]}"')
                    for r in para.runs:
                        fs = r.font.size
                        fb = r.font.bold
                        try:
                            fc = r.font.color.rgb
                        except:
                            fc = None
                        print(f'      Run: "{r.text[:80]}" size={fs} bold={fb} color={fc}')
        if shape.has_table:
            t = shape.table
            print(f'    [TABLE {len(t.rows)}x{len(t.columns)}]')
        if hasattr(shape, 'image'):
            try:
                print(f'    [IMAGE: {shape.image.content_type}]')
            except:
                pass
    print()
